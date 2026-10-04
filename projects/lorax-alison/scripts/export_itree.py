"""Build an i-Tree Eco v6 "Complete Inventory" import workbook from public/data/trees.csv.

Run from the project root (needs openpyxl: pip install openpyxl):
    python3 scripts/export_itree.py
Format follows the Eco Guide to Importing an Existing Inventory (12.14.2021):
https://www.itreetools.org/resources/manuals/Ecov6_ManualsGuides/Ecov6Guide_InventoryImporter.pdf
"""

import csv
import re
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

SRC = "public/data/trees.csv"
OUT = "exports/itree-eco/oak-park-trees-itree-eco.xlsx"
MAX_DBH = 80  # matches the app: larger recorded trunks are treated as likely data errors
SPELLING = {"Syringia": "Syringa", "Circidiphyllum": "Cercidiphyllum"}


def eco_species(latin, genus):
    """Scientific name for Eco, plus a note on how it was derived."""
    notes = []
    name = latin
    for wrong, right in SPELLING.items():
        if name.startswith(wrong) or (not name and genus == wrong):
            notes.append(f"spelling corrected from {wrong}")
    genus = SPELLING.get(genus, genus)
    for wrong, right in SPELLING.items():
        name = name.replace(wrong, right)
    if not name:
        return genus, notes + ["no Latin name in inventory; genus only"]
    if re.search(r"['\"]", name):
        name = re.sub(r"\s*['\"].*$", "", name)
        notes.append("cultivar name removed")
    words = name.split()
    if len(words) == 1 or words[-1] == "spp":
        return words[0], notes + ["genus only (species not identified)"]
    return " ".join(words), notes


def number(v):
    if v in ("", None):
        return None
    f = float(v)
    return int(f) if f.is_integer() else f


rows = list(csv.DictReader(open(SRC)))
trees, excluded = [], []
species_info = defaultdict(lambda: {"count": 0, "common": Counter(), "recorded": Counter(), "notes": set()})

for r in rows:
    dbh = number(r["dbh_in"])
    reason = None
    if r["genus"] in ("", "STUMP", "UNKNOWN"):
        reason = "stump or unknown species"
    elif dbh is None:
        reason = "no trunk diameter recorded"
    elif dbh < 0.5:
        reason = "trunk diameter under Eco's 0.5 inch minimum"
    elif dbh > MAX_DBH:
        reason = f"trunk diameter over {MAX_DBH} inches, likely a data error"
    if reason:
        excluded.append((int(r["id"]), r["common_name"], r["latin_name"], r["dbh_in"], reason))
        continue
    sp, notes = eco_species(r["latin_name"], r["genus"])
    info = species_info[sp]
    info["count"] += 1
    info["common"][r["common_name"]] += 1
    info["recorded"][r["latin_name"] or "(blank)"] += 1
    info["notes"].update(notes)
    height = number(r["height_ft"]) or None  # 0 means not recorded
    spread = number(r["spread_ft"]) or None
    trees.append([
        int(r["id"]), sp, dbh, height, spread, float(r["lat"]), float(r["lon"]), "Y", r["zone"],
        r["common_name"], r["latin_name"], r["address"], r["address_status"], r["block"],
    ])

wb = Workbook()
bold = Font(bold=True)
ref_fill = PatternFill("solid", fgColor="EEEEEE")


def sheet(ws, header, data, widths, ref_from=None):
    ws.append(header)
    for c in ws[1]:
        c.font = bold
    for row in data:
        ws.append(row)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    if ref_from:
        for col in range(ref_from, len(header) + 1):
            ws.cell(1, col).fill = ref_fill
    ws.freeze_panes = "A2"


readme = wb.active
readme.title = "Read me"
lines = [
    "Oak Park public tree inventory, formatted for i-Tree Eco v6 (Complete Inventory project)",
    "",
    f"Trees sheet: {len(trees):,} trees to import. Not included sheet: {len(excluded)} records left out, with reasons.",
    "Source: Village of Oak Park Tree Inventory (cached copy, see public/data/SOURCES.md). Built by scripts/export_itree.py.",
    "",
    "Steps (from the Eco Guide to Importing an Existing Inventory):",
    "1. In i-Tree Eco v6 (Windows): File > New Project > project type 'Complete Inventory'. Location: Oak Park, Cook County, Illinois.",
    "2. Optional: under Project Configuration, add strata named A1 to E3 to get results by the app's 15 zones.",
    "3. Data tab > Trees > Import. Choose this file and the 'Trees' sheet. Tick 'first row contains column headers'.",
    "4. Match columns: Tree ID > Tree ID; Species > Species with field type 'Scientific name'; DBH (in) > DBH;",
    "   Total tree height (ft) > Total tree height; Crown width (ft) > Crown width; Latitude and Longitude;",
    "   Public tree > Public/private; Zone > Strata (only if you did step 2).",
    "5. Skip the gray reference columns (common name, recorded Latin name, address, block). They are for joining results back.",
    "6. If some species are not matched automatically, map them by hand. The 'Species list' sheet shows each name, its count and how it was derived.",
    "7. Reports tab: submit for processing, then export the per-tree results (keep Tree ID) and add them to the project workspace.",
    "",
    "Notes on the data:",
    "- Species use the inventory's Latin name. Cultivar names were removed; genus-only names (and 'spp') are given as the genus.",
    "- Trees with no Latin name are given as their genus, taken from the inventory's common name.",
    "- Two misspelled genera were corrected: Syringia > Syringa, Circidiphyllum > Cercidiphyllum.",
    "- Heights are coded in 5 or 10 ft steps in the source. A height or crown width of 0 was treated as not recorded.",
    "- All trees are public (Y). The inventory does not say which are street trees and which are in parks, so that field is left out.",
    "- The inventory has no condition or dieback field, so all trees are treated as alive. No land use is given.",
    "- Addresses are estimates; only address_status 'match' agrees with a Cook County Assessor property address.",
]
for line in lines:
    readme.append([line])
readme["A1"].font = Font(bold=True, size=13)
readme.column_dimensions["A"].width = 130

sheet(
    wb.create_sheet("Trees"),
    ["Tree ID", "Species", "DBH (in)", "Total tree height (ft)", "Crown width (ft)", "Latitude", "Longitude",
     "Public tree", "Zone", "Common name (inventory)", "Latin name (as recorded)", "Address (estimated)",
     "Address status", "Block"],
    trees,
    [10, 28, 10, 12, 12, 12, 12, 10, 8, 26, 34, 36, 12, 24],
    ref_from=10,
)

species_rows = []
for sp, info in sorted(species_info.items(), key=lambda kv: -kv[1]["count"]):
    species_rows.append([
        sp, info["count"], info["common"].most_common(1)[0][0],
        "; ".join(f"{k} ({v})" for k, v in info["recorded"].most_common()),
        "; ".join(sorted(info["notes"])) or "as recorded",
    ])
sheet(wb.create_sheet("Species list"),
      ["Species (for Eco)", "Trees", "Most common name in inventory", "Latin names as recorded (count)", "How derived"],
      species_rows, [30, 8, 28, 70, 50])

sheet(wb.create_sheet("Not included"), ["Tree ID", "Common name", "Latin name", "DBH (in)", "Reason"],
      excluded, [10, 24, 28, 10, 50])

wb.save(OUT)
print(f"{len(trees)} trees, {len(species_rows)} species values, {len(excluded)} excluded -> {OUT}")
