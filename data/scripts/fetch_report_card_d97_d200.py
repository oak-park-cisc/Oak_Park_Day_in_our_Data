#!/usr/bin/env python3
"""
Build data/report-card-d97-d200.csv from the Illinois State Board of Education
(ISBE) Illinois Report Card public data sets.

Source: ISBE Illinois Report Card Data Library
    https://www.isbe.net/Pages/Illinois-State-Report-Card-Data.aspx
Each yearly "Report Card Public Data Set" workbook is linked from that page as
    https://www.isbe.net/_layouts/Download.aspx?SourceUrl=/Documents/<file>
which redirects to the .xlsx. Files used (report card year -> file):
    2018 Report-Card-Public-Data-Set.xlsx
    2019 2019-Report-Card-Public-Data-Set.xlsx
    2020 2020-Report-Card-Public-Data-Set.xlsx
    2021 2021-RC-Pub-Data-Set.xlsx
    2022 2022-Report-Card-Public-Data-Set.xlsx
    2023 23-RC-Pub-Data-Set.xlsx
    2024 24-RC-Pub-Data-Set.xlsx
    2025 2025-Report-Card-Public-Data-Set.xlsx
(rc-trend-data.xlsx on the same page is a statewide-only 15-year summary with
no district rows, so it is not used.)

Rows: Oak Park ESD 97 and Oak Park-River Forest SD 200 (district rows plus
each of their schools), a set of nearby comparison districts, and the
statewide row, one row per entity per report card year.

Requirements: Python 3 standard library plus `requests` and `openpyxl`.
The workbooks total about 200 MB; they are downloaded once into a cache
directory (env ISBE_CACHE_DIR, default <system temp>/isbe-report-card) and
reused on later runs. No API key is needed.

Usage:
    python3 data/scripts/fetch_report_card_d97_d200.py [output.csv]
"""

import csv
import os
import re
import sys
import tempfile

import requests

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

OUT_PATH = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
    os.path.dirname(os.path.abspath(__file__)), "..", "report-card-d97-d200.csv")
CACHE = os.environ.get("ISBE_CACHE_DIR") or os.path.join(tempfile.gettempdir(), "isbe-report-card")
BASE = "https://www.isbe.net/_layouts/Download.aspx?SourceUrl=/Documents/"

FILES = {
    2018: "Report-Card-Public-Data-Set.xlsx",
    2019: "2019-Report-Card-Public-Data-Set.xlsx",
    2020: "2020-Report-Card-Public-Data-Set.xlsx",
    2021: "2021-RC-Pub-Data-Set.xlsx",
    2022: "2022-Report-Card-Public-Data-Set.xlsx",
    2023: "23-RC-Pub-Data-Set.xlsx",
    2024: "24-RC-Pub-Data-Set.xlsx",
    2025: "2025-Report-Card-Public-Data-Set.xlsx",
}

# District RCDTS codes (15 characters; districts end in 0000). The first 11
# characters identify the district; school codes share that prefix.
DISTRICTS = {
    "060160970020000": "Oak Park ESD 97",
    "060162000130000": "Oak Park-River Forest SD 200",
    "060160900020000": "River Forest SD 90",
    "060160910020000": "Forest Park SD 91",
    "060160980020000": "Berwyn North SD 98",
    "060161000020000": "Berwyn South SD 100",
    "190222050260000": "Elmhurst CUSD 205",
    "050162020170000": "Evanston Twp HSD 202",
    "050160650040000": "Evanston CCSD 65",
    "060160960020000": "Riverside SD 96",
    "060162080170000": "Riverside-Brookfield Twp HSD 208",
    "060162090170000": "Proviso Twp HSD 209",
    "060162040170000": "Lyons Twp HSD 204",
}
SCHOOL_PREFIXES = ("06016097002", "06016200013")   # keep school rows for D97 and D200

# Output column -> candidate header names (normalized) in priority order.
# Normalization: lowercase, en/em dashes to '-', whitespace collapsed,
# trailing school-year suffixes like ' 2016-17' removed.
INDICATORS = {
    "enrollment": ["# student enrollment", "student enrollment - total"],
    "pct_low_income": ["% student enrollment - low income", "student enrollment - low income %"],
    "pct_el": ["% student enrollment - el", "student enrollment - el %"],
    "pct_iep": ["% student enrollment - iep", "student enrollment - iep %"],
    "chronic_absenteeism_pct": ["chronic absenteeism"],
    "pupil_teacher_ratio_elementary": ["pupil teacher ratio - elementary"],
    "pupil_teacher_ratio_high_school": ["pupil teacher ratio - high school"],
    "avg_class_size_all_grades": ["avg class size - all grades"],
    "instructional_expenditure_per_pupil": ["$ instructional expenditure per pupil", "instructional expenditure per pupil"],
    "operating_expenditure_per_pupil": ["$ operating expenditures", "operating expenditures"],
    "site_based_total_per_pupil_expenditure": ["$ total per-pupil expenditures - subtotal"],
    "ela_proficiency_pct": ["% ela proficiency", "ela proficiency total %"],
    "math_proficiency_pct": ["% math proficiency", "math proficiency total %"],
    "science_proficiency_pct": ["% science proficiency", "isa proficiency total %", "% isa proficiency"],
    "grad_rate_4yr_pct": ["high school 4-year graduation rate - total"],
    # 2019's workbook has only the count ("# 9th Grade on Track"), so 2019 stays blank
    "ninth_grade_on_track_pct": ["% 9th grade on track", "9th grade on track"],
    "teacher_avg_salary": ["teacher avg salary"],
    "teacher_retention_rate_pct": ["teacher retention rate"],
    "summative_designation": ["summative designation"],
}
# Sheets to scan per workbook (any sheet whose name starts with these), in
# this priority order: the first sheet that has a column wins.
SHEET_PREFIXES = ("General", "Finance", "Financial", "ELA Math Science", "ELAMathScience", "ELA and Math", "ISA")


def sheet_rank(name):
    for i, p in enumerate(SHEET_PREFIXES):
        if name.startswith(p):
            return i
    return None


def norm(name):
    s = str(name).replace("–", "-").replace("—", "-").lower()
    s = re.sub(r"\s+", " ", s).strip()
    s = re.sub(r" (19|20)\d\d-\d\d$", "", s)      # drop ' 2016-17' style suffixes
    return s


def download(fname):
    os.makedirs(CACHE, exist_ok=True)
    path = os.path.join(CACHE, fname)
    if os.path.exists(path) and os.path.getsize(path) > 100000:
        return path
    print(f"downloading {fname} ...")
    with requests.get(BASE + fname, headers={"User-Agent": "Mozilla/5.0"}, stream=True, timeout=600) as r:
        r.raise_for_status()
        with open(path + ".part", "wb") as fh:
            for chunk in r.iter_content(1 << 20):
                fh.write(chunk)
    os.replace(path + ".part", path)
    return path


def is_wanted(rcdts, level):
    if level == "state":
        return True
    if rcdts in DISTRICTS:
        return True
    return level == "school" and rcdts.startswith(SCHOOL_PREFIXES)


def classify(rec):
    """Return (rcdts, level) for a row; level is district/school/state."""
    # 2025 formats codes as 01-009-2620-26-0000; earlier years as 010092620260000
    rcdts = str(rec.get("rcdts") or "").strip().replace("-", "")
    typ = str(rec.get("type") or rec.get("level") or "").strip().lower()
    if typ.startswith("state") or rec.get("district") in ("Statewide", "State") or rcdts.startswith("65000") or (
            not rcdts and typ not in ("school", "district")):
        return "STATE", "state"
    if not typ:  # infer from the code when there is no Type/Level column
        typ = "district" if rcdts.endswith("0000") else "school"
    if typ == "school":
        return rcdts, "school"
    if typ == "district":
        return rcdts, "district"
    return rcdts, typ or "unknown"


def extract_year(year, path, unmatched):
    wb = openpyxl.load_workbook(path, read_only=True)
    entities = {}     # rcdts -> row dict
    matched_cols = {}
    sheets = [n for n in wb.sheetnames if sheet_rank(n) is not None]
    for sheet in sorted(sheets, key=lambda n: (sheet_rank(n), n)):
        ws = wb[sheet]
        rows = ws.iter_rows(values_only=True)
        header = [norm(h) if h is not None else "" for h in next(rows)]
        if "rcdts" not in header:
            continue
        # which output columns can this sheet supply?
        colmap = {}
        for out_col, candidates in INDICATORS.items():
            if out_col in matched_cols:
                continue
            for cand in candidates:
                if cand in header:
                    colmap[out_col] = header.index(cand)
                    matched_cols[out_col] = f"{sheet}: {cand}"
                    break
        idx = {h: i for i, h in enumerate(header)}
        for r in rows:
            rec = {h: r[i] for h, i in idx.items() if h}
            rcdts, level = classify(rec)
            if not is_wanted(rcdts, level):
                continue
            ent = entities.setdefault(rcdts, {
                "year": year, "rcdts": rcdts, "level": level,
                "district_name": rec.get("district") or ("State of Illinois" if level == "state" else ""),
                "school_name": rec.get("school name") or "",
                "district_type": rec.get("district type") or "",
                "school_type": rec.get("school type") or "",
                "source_file": os.path.basename(path),
            })
            for out_col, ci in colmap.items():
                val = r[ci] if ci < len(r) else None
                if val is not None and ent.get(out_col) in (None, ""):
                    ent[out_col] = val
    wb.close()
    for out_col in INDICATORS:
        if out_col not in matched_cols:
            unmatched.setdefault(year, []).append(out_col)
    print(f"{year}: {len(entities)} entities; column mapping:")
    for k, v in matched_cols.items():
        print(f"    {k:40s} <- {v}")
    return list(entities.values())


def type_label(t):
    t = str(t or "").upper()      # values vary by year: UNIT, Unit District, Elementary Distrcts, ...
    for key, label in (("ELEM", "elementary"), ("HIGH", "high"), ("UNIT", "unit")):
        if key in t:
            return label
    return t.lower()


def main():
    out, unmatched = [], {}
    for year, fname in sorted(FILES.items()):
        out.extend(extract_year(year, download(fname), unmatched))
    for row in out:
        row["type"] = type_label(row.get("district_type"))
        # blank out ISBE's placeholder values
        for k, v in list(row.items()):
            if isinstance(v, str) and v.strip() in ("", "-", "N/A", "n/a", "*", "**", "No Data"):
                row[k] = ""
    fields = ["year", "rcdts", "level", "district_name", "school_name", "type", "district_type", "school_type",
              *INDICATORS.keys(), "source_file"]
    out.sort(key=lambda r: (r["year"], r["level"] != "state", r["rcdts"]))
    with open(OUT_PATH, "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(out)
    print(f"wrote {len(out)} rows to {OUT_PATH}")
    if unmatched:
        print("indicators with no matching column, by year:")
        for y, cols in sorted(unmatched.items()):
            print(f"  {y}: {', '.join(cols)}")


if __name__ == "__main__":
    main()
